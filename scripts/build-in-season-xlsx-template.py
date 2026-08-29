"""Build the Hugo in-season workout import template (.xlsx)."""

from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "templates" / "hugo-in-season-workout-import.xlsx"

HEADERS = [
    "week",
    "day",
    "session_date",
    "focus",
    "hugo_group",
    "label",
    "name",
    "block",
    "set_count",
    "targets",
    "notes",
]

# One in-season soccer session (matches src/lib/weight-room/sample-in-season.ts)
SAMPLE_ROWS = [
    [3, "Wednesday", "2026-09-16", "Lower + power", "soccer", "W", "Warmup Circuit", "Warmup", 0, "", "Spring ankle · hip hike · lunge ISO · pogos"],
    [3, "Wednesday", "2026-09-16", "Lower + power", "soccer", "1", "CMJ (Vertical Jump)", "Primer", 3, "Max Attempt|Max Attempt|Max Attempt", "Explosive · log best (in)"],
    [3, "Wednesday", "2026-09-16", "Lower + power", "soccer", "2", "Trap-Bar Deadlift", "Main", 4, "3 @ RPE 7|3 @ RPE 8|3 @ RPE 8|3 @ RPE 8", "Stop short of grinders · in-season"],
    [3, "Wednesday", "2026-09-16", "Lower + power", "soccer", "3A", "DB Reverse Lunge", "Secondary", 3, "6/side @ RPE 7|6/side @ RPE 8|6/side @ RPE 8", "Trunk stacked"],
    [3, "Wednesday", "2026-09-16", "Lower + power", "soccer", "3B", "Copenhagen Side-Plank", "Secondary", 3, "20s/side|20s/side|20s/side", "Quality over time"],
    [3, "Wednesday", "2026-09-16", "Lower + power", "soccer", "4", "Nordic / Slider Leg Curl", "Accessory", 3, "5–6|5–6|5–6", "Eccentric control"],
    [3, "Wednesday", "2026-09-16", "Lower + power", "soccer", "5", "Med-Ball Rotation Throw", "Accessory", 2, "4/side|4/side", "Submax · hip-shoulder timing"],
]

ALLOWED_DAYS = ["Monday", "Tuesday", "Wednesday", "Thursday", "Friday"]
ALLOWED_GROUPS = [
    "soccer",
    "volleyball",
    "xc",
    "extracurricular",
    "mens_basketball",
    "womens_basketball",
    "track",
    "baseball",
]
ALLOWED_BLOCKS = [
    "Warmup",
    "Primer",
    "Main",
    "Secondary",
    "Accessory",
    "Conditioning",
]

FONT = "Arial"
NAVY = "1B365D"
GOLD = "C4A35A"
INPUT_YELLOW = "FFF3C4"
HEADER_FONT = Font(name=FONT, bold=True, color="FFFFFF", size=11)
BODY_FONT = Font(name=FONT, size=11, color="1A1A1A")
TITLE_FONT = Font(name=FONT, bold=True, size=16, color=NAVY)
SECTION_FONT = Font(name=FONT, bold=True, size=12, color=NAVY)
HINT_FONT = Font(name=FONT, size=10, italic=True, color="555555")
CODE_FONT = Font(name="Consolas", size=10, color="1A1A1A")
THIN = Border(
    left=Side(style="thin", color="D0D5DD"),
    right=Side(style="thin", color="D0D5DD"),
    top=Side(style="thin", color="D0D5DD"),
    bottom=Side(style="thin", color="D0D5DD"),
)
HEADER_FILL = PatternFill("solid", fgColor=NAVY)
INPUT_FILL = PatternFill("solid", fgColor=INPUT_YELLOW)
WRAP = Alignment(wrap_text=True, vertical="center")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)


def style_header_row(ws, row: int, cols: int) -> None:
    for col in range(1, cols + 1):
        cell = ws.cell(row, col)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = CENTER
        cell.border = THIN


def main() -> None:
    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()

    lists = wb.active
    lists.title = "Lists"
    lists["A1"] = "day"
    lists["B1"] = "hugo_group"
    lists["C1"] = "block"
    style_header_row(lists, 1, 3)
    for i, value in enumerate(ALLOWED_DAYS, start=2):
        lists.cell(i, 1, value).font = BODY_FONT
    for i, value in enumerate(ALLOWED_GROUPS, start=2):
        lists.cell(i, 2, value).font = BODY_FONT
    for i, value in enumerate(ALLOWED_BLOCKS, start=2):
        lists.cell(i, 3, value).font = BODY_FONT
    for col, width in enumerate([16, 22, 16], start=1):
        lists.column_dimensions[get_column_letter(col)].width = width
    lists.sheet_state = "hidden"

    how = wb.create_sheet("How to use", 0)
    how.sheet_view.showGridLines = False
    how.page_setup.fitToPage = True
    how.page_setup.fitToWidth = 1
    how.page_setup.fitToHeight = 1
    how.column_dimensions["A"].width = 92
    how.merge_cells("A1:A1")
    how["A1"] = "Hugo in-season workout import"
    how["A1"].font = TITLE_FONT
    how["A2"] = (
        "Copy this workbook for each sport week. Keep the Import sheet header row "
        "exactly as written — the app rejects any other first line."
    )
    how["A2"].font = HINT_FONT
    how["A2"].alignment = WRAP
    how.row_dimensions[2].height = 36

    steps = [
        ("1. Duplicate this file", "File → Save As. Name it by sport and week, e.g. soccer-week-3.xlsx."),
        (
            "2. Edit the Import sheet only",
            "One row per movement. Rows that share hugo_group + session_date + focus become one printed card. "
            "Leave Warmup set_count at 0 and targets blank. Insert new rows inside the table (Tab in the last cell) "
            "so dropdowns copy down. Delete unused rows before you export CSV — blank rows will fail import.",
        ),
        (
            "3. Dates as YYYY-MM-DD text",
            "Type 2026-09-16 (not 9/16/2026). The session_date column is formatted as text so Excel will not rewrite it.",
        ),
        (
            "4. Targets must match set_count",
            "Pipe-separate one target per set. Four sets → 3 @ RPE 7|3 @ RPE 8|3 @ RPE 8|3 @ RPE 8. "
            "A comma inside a note is fine; put quotes around the cell if you also need a comma in targets.",
        ),
        (
            "5. Stay scan-safe",
            "At most 12 movements and 6 set columns per card. Pair supersets as 3A / 3B on two rows with the same date and focus.",
        ),
        (
            "6. Save as CSV, then import",
            "In Excel: File → Save As → CSV UTF-8 (Comma delimited). On Cards / templates, choose that file. "
            "If import says invalid header, open the CSV in a text editor and confirm line 1 is exactly the header below.",
        ),
        (
            "Required header (do not rename)",
            "week,day,session_date,focus,hugo_group,label,name,block,set_count,targets,notes",
        ),
        (
            "Allowed hugo_group values",
            "soccer · volleyball · xc · extracurricular · mens_basketball · womens_basketball · track · baseball",
        ),
    ]
    row = 4
    for title, body in steps:
        how.cell(row, 1, title).font = SECTION_FONT
        how.cell(row + 1, 1, body).font = (
            CODE_FONT if title.startswith("Required") else BODY_FONT
        )
        how.cell(row + 1, 1).alignment = WRAP
        how.row_dimensions[row + 1].height = 48
        row += 3

    how["A28"] = (
        "The Import sheet includes a sample soccer Wednesday (Week 3 — Lower + power). "
        "Delete those rows after you copy the layout, or import them as a trial card."
    )
    how["A28"].font = HINT_FONT
    how["A28"].alignment = WRAP
    how.row_dimensions[28].height = 36

    ws = wb.create_sheet("Import", 0)
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A2"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.sheet_properties.tabColor = GOLD

    for col, header in enumerate(HEADERS, start=1):
        cell = ws.cell(1, col, header)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = CENTER
        cell.border = THIN

    comments = {
        1: "Integer week number, or leave blank.",
        2: "Monday … Friday (dropdown).",
        3: "Text YYYY-MM-DD. Same date + sport + focus = one card.",
        4: "Printed focus line, e.g. Lower + power.",
        5: "Must match a dropdown value (soccer, volleyball, …).",
        6: "W, 1, 2, 3A, 3B, …",
        7: "Exercise name printed on the card.",
        8: "Warmup / Primer / Main / Secondary / Accessory / Conditioning.",
        9: "Integer. 0 for circuits with no Load×Reps cells. Max 6 for scan-safe print.",
        10: "One target per set, separated by |. Count must equal set_count.",
        11: "Coach note on the printed row. Optional.",
    }
    for col, text in comments.items():
        ws.cell(1, col).comment = Comment(text, "Hugo import")

    last_data_row = 1 + len(SAMPLE_ROWS)
    for r, values in enumerate(SAMPLE_ROWS, start=2):
        for c, value in enumerate(values, start=1):
            cell = ws.cell(r, c, value)
            cell.font = BODY_FONT
            cell.border = THIN
            cell.alignment = WRAP if c in (7, 10, 11) else Alignment(vertical="center")
            cell.fill = INPUT_FILL
            if c == 3:
                cell.number_format = "@"
            if c in (1, 9):
                cell.alignment = CENTER

    widths = {
        "A": 8,
        "B": 14,
        "C": 14,
        "D": 16,
        "E": 18,
        "F": 8,
        "G": 28,
        "H": 14,
        "I": 12,
        "J": 42,
        "K": 36,
    }
    for letter, width in widths.items():
        ws.column_dimensions[letter].width = width
    ws.row_dimensions[1].height = 22
    for r in range(2, last_data_row + 1):
        ws.row_dimensions[r].height = 22

    dv_day = DataValidation(
        type="list",
        formula1="Lists!$A$2:$A$6",
        allow_blank=True,
        showDropDown=False,
        showErrorMessage=True,
        errorTitle="Day",
        error="Use Monday through Friday.",
    )
    dv_group = DataValidation(
        type="list",
        formula1="Lists!$B$2:$B$9",
        allow_blank=True,
        showDropDown=False,
        showErrorMessage=True,
        errorTitle="hugo_group",
        error="Pick a value from the list (lowercase, underscored).",
    )
    dv_block = DataValidation(
        type="list",
        formula1="Lists!$C$2:$C$7",
        allow_blank=True,
        showDropDown=False,
        showErrorMessage=True,
        errorTitle="block",
        error="Pick a training block from the list.",
    )
    dv_sets = DataValidation(
        type="whole",
        operator="between",
        formula1="0",
        formula2="6",
        allow_blank=True,
        showErrorMessage=True,
        errorTitle="set_count",
        error="Use an integer from 0 to 6 (scan-safe set columns).",
    )
    for dv, col in (
        (dv_day, "B"),
        (dv_group, "E"),
        (dv_block, "H"),
        (dv_sets, "I"),
    ):
        dv.add(f"{col}2:{col}{last_data_row}")
        ws.add_data_validation(dv)

    table = Table(displayName="WorkoutImport", ref=f"A1:K{last_data_row}")
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    ws.add_table(table)

    wb.properties.title = "Hugo in-season workout import"
    wb.properties.creator = "The LCA Speed Journal"
    wb.active = ws
    wb.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
