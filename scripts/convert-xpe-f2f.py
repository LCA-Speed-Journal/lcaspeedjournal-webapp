"""Convert Villani-Metrics_40yd-Predictors.xlsx → src/lib/norms/f2f/tables/*.json"""
from pathlib import Path
import json
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
DEFAULT_XLSX = Path(
    r"C:\Users\rossp\OneDrive\Documents\Python Scripts\2026 Coding\Villani-Metrics_40yd-Predictors.xlsx"
)
OUT = ROOT / "src" / "lib" / "norms" / "f2f" / "tables"

SHEETS = {
    "Broad Jump (Explosion)": {
        "file": "broad-jump.json",
        "quality": "explosion",
        "kind": "broad",
        "journal_split": None,
        "fly_yards": None,
    },
    "5-10 Fly (Force)": {
        "file": "fly-5-15.json",
        "quality": "force",
        "kind": "fly",
        "journal_split": "5-15yd",
        "fly_yards": 10,
    },
    "20-20 Fly (Top-End or Form)": {
        "file": "fly-20-40.json",
        "quality": "form",
        "kind": "fly",
        "journal_split": "20-40yd",
        "fly_yards": 20,
    },
    "20-10 Fly (Top-End or Form)": {
        "file": "fly-20-30.json",
        "quality": "form",
        "kind": "fly",
        "journal_split": "20-30yd",
        "fly_yards": 10,
    },
    "30-10 Fly (Top-End or Form)": {
        "file": "fly-30-40.json",
        "quality": "form",
        "kind": "fly",
        "journal_split": "30-40yd",
        "fly_yards": 10,
    },
}


def main(xlsx: Path = DEFAULT_XLSX) -> None:
    wb = load_workbook(xlsx, data_only=True, read_only=True)
    OUT.mkdir(parents=True, exist_ok=True)
    for sheet_name, spec in SHEETS.items():
        rows_out = []
        for raw in wb[sheet_name].iter_rows(min_row=2, values_only=True):
            vals = [c for c in raw if c is not None]
            if spec["kind"] == "broad":
                if len(vals) < 2:
                    continue
                rows_out.append({
                    "distance_ft": float(vals[0]),
                    "predicted_40": float(vals[1]),
                })
            else:
                if len(vals) < 3:
                    continue
                rows_out.append({
                    "time_s": float(vals[0]),
                    "mph": float(vals[1]),
                    "predicted_40": float(vals[2]),
                })
        payload = {
            "meta": {
                "source": "Tony Villani XPE Game-Speed",
                "workbook": xlsx.name,
                "sheet": sheet_name,
                "quality": spec["quality"],
                "journal_split": spec["journal_split"],
                "fly_yards": spec["fly_yards"],
                "population": "male football ~175lb+",
            },
            "rows": rows_out,
        }
        (OUT / spec["file"]).write_text(json.dumps(payload, indent=2) + "\n")
        print(spec["file"], len(rows_out))


if __name__ == "__main__":
    main()
